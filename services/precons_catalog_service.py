"""Servicio y motor de búsqueda de alta velocidad para el Catálogo Oficial PRECONS III.

Gestiona la base de datos indexada SQLite (con FTS5) que contiene los 15,981 renglones,
4,383 recursos (materiales, equipos y mano de obra), parámetros y límites oficiales
de la construcción en Cuba (RoPres 3.30).
"""

from __future__ import annotations

import os
import sqlite3
from typing import Any, Dict, List, Optional, Tuple

DEFAULT_DB_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "resources", "catalogs", "precons_iii.db"
)
DEFAULT_EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "Catalogo_PRECONS_III_Completo.xlsx"
)


class PreconsCatalogService:
    """Acceso y consultas instantáneas sobre el catálogo PRECONS III."""

    def __init__(self, db_path: Optional[str] = None):
        self.db_path = db_path or DEFAULT_DB_PATH
        if not os.path.exists(self.db_path) and os.path.exists(DEFAULT_EXCEL_PATH):
            try:
                self.build_catalog_db(DEFAULT_EXCEL_PATH, self.db_path)
            except Exception as e:
                import logging
                logging.getLogger("PreconsCatalogService").warning(
                    "No se pudo compilar la base de datos PRECONS III automáticamente: %s", e
                )

    def _conn(self) -> sqlite3.Connection:
        if not os.path.exists(self.db_path):
            raise FileNotFoundError(
                f"Base de datos de catálogo PRECONS III no encontrada en {self.db_path}. "
                "Ejecute 'precons catalogo-build' primero.")
        con = sqlite3.connect(self.db_path)
        con.row_factory = sqlite3.Row
        return con

    @staticmethod
    def build_catalog_db(excel_path: str, db_path: Optional[str] = None) -> Dict[str, Any]:
        """Compila Catalogo_PRECONS_III_Completo.xlsx en SQLite con índices FTS5."""
        import openpyxl
        out_db = db_path or DEFAULT_DB_PATH
        os.makedirs(os.path.dirname(out_db), exist_ok=True)
        if os.path.exists(out_db):
            os.remove(out_db)

        con = sqlite3.connect(out_db)
        cur = con.cursor()

        cur.execute("""
        CREATE TABLE IF NOT EXISTS renglones (
            codigo TEXT PRIMARY KEY,
            descripcion TEXT,
            unidad TEXT,
            seccion_cod TEXT,
            seccion TEXT,
            capitulo_cod TEXT,
            capitulo TEXT,
            subcapitulo_cod TEXT,
            subcapitulo TEXT,
            materiales_cup REAL,
            mano_obra_cup REAL,
            equipos_cup REAL,
            total_cup REAL
        );
        """)

        cur.execute("""
        CREATE TABLE IF NOT EXISTS recursos (
            codigo TEXT PRIMARY KEY,
            descripcion TEXT,
            tipo TEXT,
            unidad TEXT,
            precio_cup REAL,
            peso_kg REAL
        );
        """)

        cur.execute("""
        CREATE TABLE IF NOT EXISTS parametros (
            codigo TEXT PRIMARY KEY,
            nombre TEXT,
            valor REAL,
            unidad TEXT,
            nota TEXT
        );
        """)

        cur.execute("""
        CREATE VIRTUAL TABLE IF NOT EXISTS renglones_fts USING fts5(
            codigo, descripcion, seccion, capitulo, subcapitulo, content='renglones', content_rowid='rowid'
        );
        """)

        cur.execute("""
        CREATE VIRTUAL TABLE IF NOT EXISTS recursos_fts USING fts5(
            codigo, descripcion, tipo, content='recursos', content_rowid='rowid'
        );
        """)

        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

        # 1. Renglones
        renglones_count = 0
        if "RENGLONES_PRECONS" in wb.sheetnames:
            ws = wb["RENGLONES_PRECONS"]
            rows = ws.iter_rows(values_only=True)
            next(rows)  # header
            batch: List[Tuple[Any, ...]] = []
            for r in rows:
                if not r or not r[6]:
                    continue
                sec_cod, sec, cap_cod, cap, subcap_cod, subcap, cod, u, mat, mo, eq, tot, desc = r[:13]
                batch.append((
                    str(cod).strip(), str(desc or "").strip(), str(u or "").strip(),
                    str(sec_cod or "").strip(), str(sec or "").strip(),
                    str(cap_cod or "").strip(), str(cap or "").strip(),
                    str(subcap_cod or "").strip(), str(subcap or "").strip(),
                    float(mat or 0.0), float(mo or 0.0), float(eq or 0.0), float(tot or 0.0)
                ))
            cur.executemany("INSERT OR REPLACE INTO renglones VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", batch)
            renglones_count = len(batch)

        # 2. Materiales, Equipos, Mano de obra
        recursos_batch: List[Tuple[Any, ...]] = []
        if "MATERIALES" in wb.sheetnames:
            ws = wb["MATERIALES"]
            rows = ws.iter_rows(values_only=True)
            next(rows)
            for r in rows:
                if not r or not r[0]:
                    continue
                cod, desc, u, p, w = r[:5]
                recursos_batch.append((
                    str(cod).strip(), str(desc or "").strip(), "MATERIAL",
                    str(u or "").strip(), float(p or 0.0), float(w or 0.0)
                ))

        if "EQUIPOS" in wb.sheetnames:
            ws = wb["EQUIPOS"]
            rows = ws.iter_rows(values_only=True)
            next(rows)
            for r in rows:
                if not r or not r[0]:
                    continue
                cod, desc, u, p = r[:4]
                recursos_batch.append((
                    str(cod).strip(), str(desc or "").strip(), "EQUIPMENT",
                    str(u or "").strip(), float(p or 0.0), 0.0
                ))

        if "MANO_OBRA" in wb.sheetnames:
            ws = wb["MANO_OBRA"]
            rows = ws.iter_rows(values_only=True)
            next(rows)
            for r in rows:
                if not r or not r[0]:
                    continue
                cod, desc, u, p = r[:4]
                recursos_batch.append((
                    str(cod).strip(), str(desc or "").strip(), "LABOR",
                    str(u or "").strip(), float(p or 0.0), 0.0
                ))

        cur.executemany("INSERT OR REPLACE INTO recursos VALUES (?,?,?,?,?,?)", recursos_batch)

        # 3. Parametros y Limites
        param_count = 0
        for sheet in ["PARAMETROS", "LIMITES"]:
            if sheet in wb.sheetnames:
                ws = wb[sheet]
                rows = ws.iter_rows(values_only=True)
                next(rows)
                for r in rows:
                    if not r or not r[0]:
                        continue
                    cod, nom, val, u = r[:4]
                    nota = str(r[4]) if len(r) > 4 and r[4] is not None else ""
                    try:
                        val_num = float(val or 0.0)
                    except Exception:
                        val_num = 0.0
                    cur.execute("INSERT OR REPLACE INTO parametros VALUES (?,?,?,?,?)",
                                (str(cod).strip(), str(nom or "").strip(), val_num, str(u or "").strip(), nota))
                    param_count += 1

        # Reconstruir indices de texto completo
        cur.execute("INSERT INTO renglones_fts(renglones_fts) VALUES('rebuild')")
        cur.execute("INSERT INTO recursos_fts(recursos_fts) VALUES('rebuild')")

        con.commit()
        con.close()
        wb.close()

        return {
            "db_path": out_db,
            "renglones": renglones_count,
            "recursos": len(recursos_batch),
            "parametros": param_count,
        }

    def search_renglones(self, query: str, limit: int = 50) -> List[Dict[str, Any]]:
        """Busca renglones por código o texto en < 1 milisegundo."""
        con = self._conn()
        try:
            cur = con.cursor()
            cleaned_words = [w.strip() for w in query.replace("'", " ").replace('"', " ").split() if len(w.strip()) > 1]
            if not cleaned_words:
                cur.execute("SELECT * FROM renglones LIMIT ?", (limit,))
                return [dict(r) for r in cur.fetchall()]

            fts_expr = " AND ".join(f'"{w}"*' for w in cleaned_words)
            sql = """
            SELECT r.* FROM renglones_fts f
            JOIN renglones r ON r.rowid = f.rowid
            WHERE renglones_fts MATCH ?
            ORDER BY rank
            LIMIT ?;
            """
            cur.execute(sql, (fts_expr, limit))
            rows = cur.fetchall()
            if not rows:
                # Fallback a LIKE para consultas parciales cortas
                like_term = f"%{query.strip()}%"
                cur.execute(
                    "SELECT * FROM renglones WHERE codigo LIKE ? OR descripcion LIKE ? LIMIT ?",
                    (like_term, like_term, limit)
                )
                rows = cur.fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()

    def get_renglon(self, code: str) -> Optional[Dict[str, Any]]:
        """Devuelve un renglón por su código PRECONS exacto."""
        con = self._conn()
        try:
            cur = con.cursor()
            cur.execute("SELECT * FROM renglones WHERE codigo = ?", (str(code).strip(),))
            row = cur.fetchone()
            return dict(row) if row else None
        finally:
            con.close()

    def search_recursos(self, query: str, kind: Optional[str] = None, limit: int = 50) -> List[Dict[str, Any]]:
        """Busca recursos (materiales, equipos, mano de obra) con filtro opcional."""
        con = self._conn()
        try:
            cur = con.cursor()
            cleaned_words = [w.strip() for w in query.replace("'", " ").replace('"', " ").split() if len(w.strip()) > 1]
            params: List[Any] = []
            if not cleaned_words:
                if kind:
                    cur.execute("SELECT * FROM recursos WHERE tipo = ? LIMIT ?", (kind.upper(), limit))
                else:
                    cur.execute("SELECT * FROM recursos LIMIT ?", (limit,))
                return [dict(r) for r in cur.fetchall()]

            fts_expr = " AND ".join(f'"{w}"*' for w in cleaned_words)
            if kind:
                sql = """
                SELECT r.* FROM recursos_fts f
                JOIN recursos r ON r.rowid = f.rowid
                WHERE recursos_fts MATCH ? AND r.tipo = ?
                ORDER BY rank
                LIMIT ?;
                """
                params = [fts_expr, kind.upper(), limit]
            else:
                sql = """
                SELECT r.* FROM recursos_fts f
                JOIN recursos r ON r.rowid = f.rowid
                WHERE recursos_fts MATCH ?
                ORDER BY rank
                LIMIT ?;
                """
                params = [fts_expr, limit]
            cur.execute(sql, params)
            return [dict(r) for r in cur.fetchall()]
        finally:
            con.close()

    def get_parametros(self) -> Dict[str, Any]:
        """Devuelve parámetros y límites oficiales."""
        con = self._conn()
        try:
            cur = con.cursor()
            cur.execute("SELECT * FROM parametros")
            return {r["codigo"]: {"nombre": r["nombre"], "valor": r["valor"], "unidad": r["unidad"], "nota": r["nota"]} for r in cur.fetchall()}
        finally:
            con.close()

    def calculate_apu(self, code: str, quantity: float = 1.0,
                      indirect_pct: Optional[float] = None,
                      benefit_pct: Optional[float] = None,
                      transport_pct: Optional[float] = None) -> Dict[str, Any]:
        """Calcula el Análisis de Precio Unitario (APU) según metodología PRECONS III."""
        item = self.get_renglon(code)
        if not item:
            raise ValueError(f"Renglón PRECONS {code} no encontrado")

        params = self.get_parametros()
        ind_pct = indirect_pct if indirect_pct is not None else 18.0
        ben_pct = benefit_pct if benefit_pct is not None else (params.get("UTILIDAD_MAX", {}).get("valor", 0.15) * 100.0)
        tra_pct = transport_pct if transport_pct is not None else 2.5

        mat = float(item["materiales_cup"])
        mo = float(item["mano_obra_cup"])
        eq = float(item["equipos_cup"])
        direct_cost = mat + mo + eq

        transport = direct_cost * (tra_pct / 100.0)
        direct_with_transport = direct_cost + transport
        indirect = direct_with_transport * (ind_pct / 100.0)
        benefit = (direct_with_transport + indirect) * (ben_pct / 100.0)
        unit_price = round(direct_with_transport + indirect + benefit, 2)
        total = round(unit_price * quantity, 2)

        return {
            "codigo": item["codigo"],
            "descripcion": item["descripcion"],
            "unidad": item["unidad"],
            "seccion": item["seccion"],
            "cantidad": quantity,
            "costo_directo": {
                "materiales": mat,
                "mano_obra": mo,
                "equipos": eq,
                "subtotal": direct_cost,
            },
            "coeficientes": {
                "transporte_pct": tra_pct,
                "transporte_cup": round(transport, 2),
                "indirectos_pct": ind_pct,
                "indirectos_cup": round(indirect, 2),
                "beneficio_pct": ben_pct,
                "beneficio_cup": round(benefit, 2),
            },
            "precio_unitario": unit_price,
            "total_cup": total,
            "moneda": "CUP",
        }
