"""Table file readers for the import pipeline (spec 94, PARSE stage).

Supports .xlsx/.xlsm (openpyxl) and .csv/.txt (stdlib csv). Every sheet
becomes a list of dict rows keyed by its first non-empty row (header).
Cell values keep their native type when possible; headers are stripped.
"""

from __future__ import annotations

import csv
import os
from typing import Any, Dict, List

from core.errors import ImportErrorARQ

XLSX_SUFFIXES = (".xlsx", ".xlsm")
CSV_SUFFIXES = (".csv", ".txt")
DELIMITERS = (";", ",", "\t", "|")


def is_table_file(path: str) -> bool:
    return str(path).lower().endswith(XLSX_SUFFIXES + CSV_SUFFIXES)


def _clean(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _rows_from_matrix(matrix: List[List[Any]]) -> List[Dict[str, Any]]:
    """First non-empty row is the header; the rest become dict rows."""
    rows: List[Dict[str, Any]] = []
    header: List[str] = []
    for raw in matrix:
        if not any(v is not None and str(v).strip() != "" for v in raw):
            continue
        if not header:
            header = [_clean(v) or f"col_{i}" for i, v in enumerate(raw)]
            continue
        row: Dict[str, Any] = {}
        for i, name in enumerate(header):
            row[name] = raw[i] if i < len(raw) else None
        rows.append(row)
    return rows


def _open_xlsx(path: str) -> Dict[str, Any]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportErrorARQ(
            message=f"No se pudo abrir el libro Excel: {exc}",
            code="ARQ-IMP-034", context={"path": path}) from exc
    sheets: Dict[str, List[Dict[str, Any]]] = {}
    try:
        for name in wb.sheetnames:
            matrix = [list(row) for row in wb[name].iter_rows(values_only=True)]
            sheets[name] = _rows_from_matrix(matrix)
    finally:
        wb.close()
    return {"format": "xlsx", "sheets": sheets}


def _pick_delimiter(sample_lines: List[List[str]]) -> str:
    """Deterministic delimiter choice: the one most frequent in the header."""
    first = next((row for row in sample_lines if any(str(c).strip() for c in row)), [])
    counts = {d: sum(str(c).count(d) for c in first) for d in DELIMITERS}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] > 0 else ","


def _open_csv(path: str) -> Dict[str, Any]:
    text_rows: List[List[str]] = []
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            with open(path, "r", encoding=encoding, newline="") as fh:
                for row in csv.reader(fh):
                    text_rows.append(row)
            break
        except UnicodeDecodeError:
            text_rows = []
            continue
        except OSError as exc:
            raise ImportErrorARQ(
                message=f"No se pudo leer el CSV: {exc}",
                code="ARQ-IMP-035", context={"path": path}) from exc
    if not text_rows:
        raise ImportErrorARQ(
            message="El archivo CSV está vacío o no se pudo decodificar",
            code="ARQ-IMP-031", context={"path": path})
    delimiter = _pick_delimiter(text_rows[:5])
    if delimiter != ",":
        merged = ["\u241f".join(row) for row in text_rows]
        text_rows = [list(csv.reader([line], delimiter=delimiter))[0] for line in merged]
    return {"format": "csv", "sheets": {"csv": _rows_from_matrix(text_rows)}}


def read_table(path: str) -> Dict[str, Any]:
    """PARSE stage: file -> {"format": "xlsx"|"csv", "sheets": {name: [rows]}}."""
    if not os.path.isfile(path):
        raise ImportErrorARQ(
            message=f"Archivo de importación inexistente: {path}",
            code="ARQ-IMP-030", context={"path": path})
    lower = path.lower()
    if lower.endswith(XLSX_SUFFIXES):
        return _open_xlsx(path)
    if lower.endswith(CSV_SUFFIXES):
        return _open_csv(path)
    raise ImportErrorARQ(
        message=f"Formato de tabla no soportado: {os.path.basename(path)}",
        code="ARQ-IMP-033",
        suggested_action="Use .xlsx, .xlsm, .csv o .txt delimitado.")


def to_numeric(value: Any, default: float = 0.0) -> float:
    """Excel/CSV numbers may arrive as text with comma decimal separator."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


__all__ = ["read_table", "is_table_file", "to_numeric",
           "XLSX_SUFFIXES", "CSV_SUFFIXES"]
