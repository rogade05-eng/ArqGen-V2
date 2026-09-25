"""Import pipeline tests (spec 94): tables, catalog and DXF importers.

Covers the 9-stage pipeline for the three supported sources:
- .xlsx/.csv tabular bases -> PRECONS ruleset (spec 59-60),
- JSON external bases through precons import (regression),
- DXF drawings -> wall/space entities inside a live project.
"""

from __future__ import annotations

import json
import os
import unittest

from tests.base import ARQGenTestCase

from core.errors import ImportErrorARQ
from importers.catalog import CatalogImporter
from importers.dxf_reader import DxfImporter
from importers.tables import read_table, to_numeric


class TablesTest(ARQGenTestCase):
    def _make_xlsx(self, name: str, sheets: dict) -> str:
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        for sheet, rows in sheets.items():
            ws = wb.create_sheet(sheet)
            for row in rows:
                ws.append(row)
        path = os.path.join(self.tmp, name)
        wb.save(path)
        return path

    def test_01_xlsx_headers_and_rows(self):
        path = self._make_xlsx("tabla.xlsx", {
            "PARTIDAS": [["Código", "Descripción", "UM"],
                         ["E-01", "Muro de bloque", "m3"],
                         ["E-04", "Hormigón", "m3"]],
        })
        parsed = read_table(path)
        self.assertEqual(parsed["format"], "xlsx")
        rows = parsed["sheets"]["PARTIDAS"]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["Código"], "E-01")
        self.assertEqual(rows[1]["Descripción"], "Hormigón")

    def test_02_csv_semicolon_and_empty_rows(self):
        path = os.path.join(self.tmp, "tabla.csv")
        with open(path, "w", encoding="utf-8-sig") as fh:
            fh.write("partida;nombre;unidad\r\n")
            fh.write("E-01;Muro;m3\r\n")
            fh.write(";;\r\n")
            fh.write("E-04;Hormigón;m3\r\n")
        parsed = read_table(path)
        rows = parsed["sheets"]["csv"]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["partida"], "E-01")

    def test_03_unsupported_and_missing(self):
        with self.assertRaises(ImportErrorARQ):
            read_table(os.path.join(self.tmp, "no_existente.xlsx"))
        fake = os.path.join(self.tmp, "datos.pdf")
        with open(fake, "w") as fh:
            fh.write("x")
        with self.assertRaises(ImportErrorARQ):
            read_table(fake)

    def test_04_to_numeric_locales(self):
        self.assertEqual(to_numeric("1.234,56"), 1234.56)
        self.assertEqual(to_numeric("12.5"), 12.5)
        self.assertEqual(to_numeric("12,5"), 12.5)
        self.assertEqual(to_numeric(""), 0.0)
        self.assertEqual(to_numeric(None, 2.0), 2.0)
        self.assertEqual(to_numeric(7), 7.0)


class CatalogFlatTest(ARQGenTestCase):
    """Flat shape: one row per (partida, insumo)."""

    def _make_xlsx(self, name: str, rows: list) -> str:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "BASE"
        for row in rows:
            ws.append(row)
        path = os.path.join(self.tmp, name)
        wb.save(path)
        return path

    def test_05_flat_import_and_out(self):
        path = self._make_xlsx("catalogo.xlsx", [
            ["Partida", "Nombre", "UM", "Recurso", "Rendimiento", "Merma"],
            ["E-01", "Muro de bloque", "m3", "BLOQUE", 72, 5],
            ["E-01", "", "", "CEMENTO", 18, 3],
            ["E-01", "", "", "MO-ALB", 1.6, 0],
            ["E-04", "Hormigón 210", "m3", "CEMENTO", 7.8, 2],
            ["", "fila sin código", "", "", "", ""],
        ])
        importer = CatalogImporter()
        out = os.path.join(self.tmp, "ruleset.json")
        payload, report = importer.apply(path, code="PRECONS_III",
                                         name="Catálogo completo",
                                         out_path=out)
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["shape"], "flat")
        self.assertEqual(report["counts"]["work_items"], 2)
        self.assertEqual(report["counts"]["indicators"], 4)
        self.assertEqual(report["counts"]["skipped_rows"], 1)
        self.assertEqual(payload["code"], "PRECONS_III")
        e01 = next(i for i in payload["work_items"] if i["code"] == "E-01")
        self.assertEqual(e01["name"], "Muro de bloque")
        self.assertEqual(e01["unit"], "m3")
        self.assertEqual(len(e01["indicators"]), 3)
        self.assertEqual(e01["indicators"][0]["resource"], "BLOQUE")
        self.assertEqual(e01["indicators"][0]["yield"], 72.0)
        # written file reloads with the versioned ruleset loader
        from engines.precons_engine import PreconsRuleset
        ruleset = PreconsRuleset.from_file(out)
        self.assertEqual(ruleset.code, "PRECONS_III")
        self.assertEqual(len(ruleset.work_items), 2)

    def test_06_preview_does_not_write(self):
        path = self._make_xlsx("catalogo2.xlsx", [
            ["Partida", "Nombre", "UM", "Recurso", "Rendimiento", "Merma"],
            ["E-01", "Muro", "m3", "BLOQUE", 72, 5],
        ])
        importer = CatalogImporter()
        report = importer.preview(path)
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["counts"]["work_items"], 1)

    def test_07_missing_resource_column(self):
        path = self._make_xlsx("malo.xlsx", [
            ["Partida", "Nombre", "UM"], ["E-01", "Muro", "m3"]])
        importer = CatalogImporter()
        with self.assertRaises(ImportErrorARQ):
            importer.convert(path)

    def test_08_no_items(self):
        path = self._make_xlsx("vacio.xlsx", [
            ["Partida", "Recurso", "Rendimiento"], ["", "", ""]])
        importer = CatalogImporter()
        with self.assertRaises(ImportErrorARQ):
            importer.convert(path)

    def test_09_explicit_mapping_overrides(self):
        path = self._make_xlsx("raro.xlsx", [
            ["CLAVE", "TITULO", "INSUMO", "CANTIDAD", "MERMA"],
            ["E-01", "Muro", "BLOQUE", 72, 5],
        ])
        mapping = {"columns": {"item_code": "CLAVE", "item_name": "TITULO",
                               "res_code": "INSUMO", "yield": "CANTIDAD",
                               "waste": "MERMA"}}
        payload, report = CatalogImporter(mapping=mapping).apply(path)
        self.assertEqual(report["errors"], [])
        self.assertEqual(payload["work_items"][0]["code"], "E-01")
        self.assertEqual(payload["work_items"][0]["name"], "Muro")


class CatalogMultiSheetTest(ARQGenTestCase):
    def _make_xlsx(self) -> str:
        from openpyxl import Workbook
        wb = Workbook()
        items = wb.active
        items.title = "PARTIDAS"
        for row in [["Código", "Descripción", "UM"],
                    ["E-01", "Muro de bloque", "m3"],
                    ["E-04", "Hormigón 210", "m3"]]:
            items.append(row)
        wb.create_sheet("RECURSOS").append(
            ["Código", "Nombre", "Tipo", "UM"])
        for row in [["BLOQUE", "Bloque 20", "MATERIAL", "und"],
                    ["CEMENTO", "Cemento P-35", "MATERIAL", "saco"],
                    ["MO-ALB", "Albañil", "LABOR", "jornada"]]:
            wb["RECURSOS"].append(row)
        wb.create_sheet("RENDIMIENTOS").append(
            ["Partida", "Insumo", "Rendimiento", "Desperdicio"])
        for row in [["E-01", "BLOQUE", 72, 5], ["E-01", "CEMENTO", 18, 3],
                    ["E-01", "MO-ALB", 1.6, 0],
                    ["E-04", "CEMENTO", 7.8, 2],
                    ["E-99", "BLOQUE", 1, 0]]:
            wb["RENDIMIENTOS"].append(row)
        wb.create_sheet("PRECIOS").append(["Código", "Precio"])
        for row in [["BLOQUE", 12.5], ["CEMENTO", 1450], ["MO-ALB", 350]]:
            wb["PRECIOS"].append(row)
        path = os.path.join(self.tmp, "precons_iii.xlsx")
        wb.save(path)
        return path

    def test_10_multi_sheet_full_pipeline(self):
        path = self._make_xlsx()
        importer = CatalogImporter()
        _, report, prices = importer.convert(path)
        self.assertEqual(report["detected"]["items_sheet"], "PARTIDAS")
        self.assertEqual(report["detected"]["resources_sheet"], "RECURSOS")
        self.assertEqual(report["detected"]["indicators_sheet"], "RENDIMIENTOS")
        self.assertEqual(report["detected"]["prices_sheet"], "PRECIOS")
        self.assertEqual(report["counts"]["work_items"], 2)
        self.assertEqual(report["counts"]["resources"], 3)
        self.assertEqual(report["counts"]["indicators"], 4)
        self.assertEqual(prices["count"], 3)
        self.assertEqual(prices["sample"][0], {"code": "BLOQUE", "price": 12.5})
        # ghost indicator row E-99 -> warning, not error
        self.assertTrue(any("E-99" in w for w in report["warnings"]))
        self.assertEqual(report["errors"], [])

    def test_11_csv_flat_equivalent(self):
        path = os.path.join(self.tmp, "base.csv")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("Partida;Nombre;UM;Recurso;Rendimiento;Merma\n")
            fh.write("E-01;Muro;m3;BLOQUE;72;5\n")
            fh.write("E-01;;m3;CEMENTO;18;3\n")
        payload, report = CatalogImporter().apply(path, code="CSV_IMP")
        self.assertEqual(report["errors"], [])
        self.assertEqual(len(payload["work_items"][0]["indicators"]), 2)

    def test_12_json_path_regression(self):
        path = os.path.join(self.tmp, "base.json")
        with open(path, "w", encoding="utf-8") as fh:
            json.dump({"work_items": [{"code": "E-01", "name": "Muro",
                                       "unit": "m3",
                                       "indicators": [
                                           {"resource": "BLOQUE",
                                            "yield": 72.0}]}],
                       "resources": [{"code": "BLOQUE", "name": "Bloque",
                                      "kind": "MATERIAL", "unit": "und"}]}, fh)
        payload, report = CatalogImporter().apply(path)
        self.assertEqual(report["format"], "json")
        self.assertEqual(len(payload["work_items"]), 1)


class DxfImportTest(ARQGenTestCase):
    def _make_dxf(self, name: str = "plano.dxf"):
        import ezdxf
        doc = ezdxf.new("R2010")
        msp = doc.modelspace()
        doc.layers.add(name="ARQ-MURO", color=7)
        doc.layers.add(name="ARQ-LOCAL", color=8)
        # two wall axes (LINE) on the wall layer
        msp.add_line((0, 0), (4, 0), dxfattribs={"layer": "ARQ-MURO"})
        msp.add_line((0, 3), (4, 3), dxfattribs={"layer": "ARQ-MURO"})
        msp.add_line((0, 0), (0, 0), dxfattribs={"layer": "ARQ-MURO"})  # zero
        # one closed space boundary + its name text
        msp.add_lwpolyline([(1, 1), (3, 1), (3, 2.5), (1, 2.5)], close=True,
                           dxfattribs={"layer": "ARQ-LOCAL"})
        msp.add_text("SALA", dxfattribs={"layer": "ARQ-LOCAL",
                                         "insert": (1.5, 1.5)})
        # unrelated layer: must be ignored
        msp.add_line((9, 9), (10, 10), dxfattribs={"layer": "OTRA"})
        path = os.path.join(self.tmp, name)
        doc.saveas(path)
        return path

    def test_13_dxf_preview(self):
        path = self._make_dxf()
        report = DxfImporter().preview(path)
        self.assertEqual(report["counts"]["walls"], 2)
        self.assertEqual(report["counts"]["spaces"], 1)
        self.assertEqual(report["counts"]["zero_length_or_invalid"], 1)
        self.assertEqual(report["spaces_preview"][0]["name"], "SALA")
        self.assertAlmostEqual(report["spaces_preview"][0]["area_m2"], 3.0, places=3)
        self.assertIn("OTRA", report["layers_seen"])
        self.assertEqual(report["errors"], [])

    def test_14_dxf_apply_creates_entities(self):
        path = self._make_dxf()
        from services.architecture_service import ArchitectureService
        context = self.application.create_project(
            self.project_path("imp.arqgen"), name="Importación")
        architecture = ArchitectureService(context)
        counts = DxfImporter().apply(path, architecture, "PLANTA_IMPORT")
        context.commit()
        self.assertEqual(counts["walls"], 2)
        self.assertEqual(counts["spaces"], 1)
        self.assertEqual(counts["level"], "PLANTA_IMPORT")
        walls = context.architecture.list("WALL", context.project.id)
        spaces = context.architecture.list("SPACE", context.project.id)
        self.assertEqual(len(walls), 2)
        self.assertEqual(len(spaces), 1)
        self.assertEqual(spaces[0].name, "SALA")
        self.assertAlmostEqual(spaces[0].area_m2(), 3.0, places=3)
        # deterministic ordering: sorted by y then x
        self.assertEqual(sorted(round(w.start[1], 3) for w in walls), [0.0, 3.0])

    def test_15_dxf_apply_is_audited(self):
        path = self._make_dxf()
        from services.architecture_service import ArchitectureService
        context = self.application.create_project(
            self.project_path("aud.arqgen"), name="Auditoría")
        DxfImporter().apply(path, ArchitectureService(context), "N1")
        context.commit()
        rows = context.audit_repo.query(command="IMPORT_DXF",
                                        object_id=context.project.id)
        self.assertTrue(rows)

    def test_16_dxf_no_candidates_error(self):
        import ezdxf
        doc = ezdxf.new("R2010")
        doc.modelspace().add_line((0, 0), (1, 1),
                                  dxfattribs={"layer": "SIN_SENTIDO"})
        path = os.path.join(self.tmp, "vacio.dxf")
        doc.saveas(path)
        importer = DxfImporter()
        report = importer.preview(path)
        self.assertTrue(report["errors"])

    def test_17_dxf_explicit_layers(self):
        import ezdxf
        doc = ezdxf.new("R2010")
        msp = doc.modelspace()
        doc.layers.add(name="STRUCT", color=7)
        doc.layers.add(name="ROOMS", color=8)
        msp.add_line((0, 0), (2, 0), dxfattribs={"layer": "STRUCT"})
        msp.add_lwpolyline([(0, 0), (2, 0), (2, 2), (0, 2)], close=True,
                           dxfattribs={"layer": "ROOMS"})
        path = os.path.join(self.tmp, "capas.dxf")
        doc.saveas(path)
        report = DxfImporter(layer_walls="STRUCT",
                             layer_spaces="ROOMS").preview(path)
        self.assertEqual(report["counts"]["walls"], 1)
        self.assertEqual(report["counts"]["spaces"], 1)


class CliImportTest(ARQGenTestCase):
    def _make_xlsx(self) -> str:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "BASE"
        for row in [["Partida", "Nombre", "UM", "Recurso", "Rendimiento"],
                    ["E-01", "Muro", "m3", "BLOQUE", 72]]:
            ws.append(row)
        path = os.path.join(self.tmp, "cli.xlsx")
        wb.save(path)
        return path

    def test_18_importar_catalogo_cli(self):
        from app.cli import run
        path = self._make_xlsx()
        out = os.path.join(self.tmp, "salida.json")
        code = run(["importar", "catalogo", path, "--code", "CLI_IMP",
                    "--out", out])
        self.assertEqual(code, 0)
        with open(out, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
        self.assertEqual(payload["code"], "CLI_IMP")
        self.assertEqual(len(payload["work_items"]), 1)
        # preview mode (sin --out): exit 0, ningún archivo escrito
        code = run(["importar", "catalogo", path])
        self.assertEqual(code, 0)
        self.assertEqual(sorted(os.listdir(self.tmp)), ["cli.xlsx", "salida.json"])

    def test_19_importar_dxf_cli_preview_and_apply(self):
        from app.cli import run
        dxf_path = DxfImportTest._make_dxf(
            self, "cli.dxf")  # reuse fixture builder
        db = self.project_path("cli_dxf.arqgen")
        self.application.create_project(db, name="CLI DXF").close()
        # PREVIEW first (no --aplicar): no entities created
        code = run(["importar", "dxf", dxf_path, "--file", db])
        self.assertEqual(code, 0)
        context = self.application.open_project(db)
        self.assertEqual(len(context.architecture.list(
            "WALL", context.project.id)), 0)
        context.close()
        # APPLY
        code = run(["importar", "dxf", dxf_path, "--file", db,
                    "--nivel", "P1", "--aplicar"])
        self.assertEqual(code, 0)
        context = self.application.open_project(db)
        self.assertEqual(len(context.architecture.list(
            "WALL", context.project.id)), 2)
        self.assertEqual(len(context.architecture.list(
            "SPACE", context.project.id)), 1)
        context.close()

    def test_20_precons_import_accepts_xlsx(self):
        from app.cli import run
        path = self._make_xlsx()
        db = self.project_path("precons_cli.arqgen")
        self.application.create_project(db, name="Precons CLI").close()
        out = os.path.join(self.tmp, "imp.json")
        code = run(["precons", "import", db, "--data", path,
                    "--code", "XLS_IMP", "--out", out])
        self.assertEqual(code, 0)
        with open(out, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
        self.assertEqual(payload["code"], "XLS_IMP")


if __name__ == "__main__":
    unittest.main()
